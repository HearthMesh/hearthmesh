from pathlib import Path
from datetime import date
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs"
QA = OUT / "qa"
BRAND = ROOT / "brand"
OUT.mkdir(parents=True, exist_ok=True)
QA.mkdir(parents=True, exist_ok=True)

INK = "111820"
EMBER = "F47B45"
CYAN = "54C8D8"
PALE = "EAF5F6"
GRAY = "D9E0E2"
MID = "65757B"
WHITE = "FFFFFF"


def font(size, bold=False):
    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def rounded_box(draw, xy, fill, outline, radius=20, width=3):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def arrow(draw, start, end, color, width=5):
    draw.line([start, end], fill=color, width=width)
    x, y = end
    draw.polygon([(x, y), (x - 15, y - 9), (x - 15, y + 9)], fill=color)


def make_architecture_diagram(path):
    img = Image.new("RGB", (1800, 1040), "white")
    d = ImageDraw.Draw(img)
    title = font(42, True)
    body = font(27)
    small = font(22)
    d.text((80, 45), "HearthMesh target architecture", fill="#" + INK, font=title)
    d.text((80, 105), "Public continuity mesh with separately isolated private cells", fill="#" + MID, font=body)

    boxes = [
        ((90, 210, 500, 420), "Publisher tools", ["Build manifest", "Hash files", "Sign HearthPack"], EMBER),
        ((695, 180, 1105, 450), "HearthMesh node", ["Identity and peer policy", "Verification and storage", "Replication and repair", "Local read-only portal"], CYAN),
        ((1300, 210, 1710, 420), "Known peer nodes", ["Direct connections", "Verified inventories", "Pull replication"], CYAN),
        ((695, 650, 1105, 890), "Private continuity cell", ["Separate membership", "Isolated namespace", "Future implementation"], EMBER),
    ]
    for (x1, y1, x2, y2), heading, lines, accent in boxes:
        rounded_box(d, (x1, y1, x2, y2), "#F7FAFA", "#" + accent, 22, 4)
        d.rectangle((x1, y1, x2, y1 + 14), fill="#" + accent)
        d.text((x1 + 30, y1 + 38), heading, fill="#" + INK, font=font(30, True))
        yy = y1 + 96
        for line in lines:
            d.ellipse((x1 + 32, yy + 10, x1 + 43, yy + 21), fill="#" + accent)
            d.text((x1 + 58, yy), line, fill="#" + MID, font=small)
            yy += 42
    arrow(d, (500, 315), (680, 315), "#" + EMBER)
    arrow(d, (1105, 315), (1285, 315), "#" + CYAN)
    arrow(d, (900, 465), (900, 630), "#" + MID)
    d.text((535, 252), "local submit", fill="#" + MID, font=small)
    d.text((1170, 245), "peer API", fill="#" + MID, font=small)
    d.text((930, 530), "explicit boundary", fill="#" + MID, font=small)
    d.text((90, 970), "Target architecture. The MVP implements a three-node allowlisted subset and does not implement private-cell confidentiality.", fill="#" + MID, font=small)
    img.save(path, quality=95)


def make_recovery_diagram(path):
    img = Image.new("RGB", (1800, 760), "white")
    d = ImageDraw.Draw(img)
    d.text((80, 45), "MVP failure and recovery flow", fill="#" + INK, font=font(42, True))
    labels = [
        ("A", "Publisher", "offline", EMBER),
        ("B", "Replica", "corrupt", "D44D5C"),
        ("C", "Replica", "valid", CYAN),
        ("B", "Replica", "restored", CYAN),
    ]
    xs = [130, 570, 1010, 1450]
    for i, (letter, role, state, color) in enumerate(labels):
        x = xs[i]
        d.ellipse((x, 245, x + 190, 435), fill="#F7FAFA", outline="#" + color, width=6)
        d.text((x + 75, 275), letter, fill="#" + color, font=font(46, True))
        tw = d.textbbox((0, 0), role, font=font(25, True))[2]
        d.text((x + 95 - tw / 2, 345), role, fill="#" + INK, font=font(25, True))
        sw = d.textbbox((0, 0), state.upper(), font=font(19, True))[2]
        d.text((x + 95 - sw / 2, 455), state.upper(), fill="#" + color, font=font(19, True))
        if i < len(labels) - 1:
            arrow(d, (x + 215, 340), (xs[i + 1] - 25, 340), "#" + MID, 4)
    d.text((300, 265), "publisher loss", fill="#" + MID, font=font(20))
    d.text((740, 265), "digest failure", fill="#" + MID, font=font(20))
    d.text((1175, 265), "verified repair", fill="#" + MID, font=font(20))
    d.text((80, 630), "Repair succeeds only while another reachable peer retains bytes matching the signed manifest.", fill="#" + MID, font=font(24))
    img.save(path, quality=95)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + name))
        if node is None:
            node = OxmlElement("w:" + name)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_keep(paragraph, next_paragraph=False):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext" if next_paragraph else "w:keepLines")
    p_pr.append(keep)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    set_keep(p, True)
    return p


def add_para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered(doc, items):
    for number, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        p.add_run(f"{number}.").bold = True
        p.add_run("  " + item)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    prevent_row_split(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, INK)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)
        if widths:
            cell.width = Inches(widths[i])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for i, value in enumerate(row):
            cell = cells[i]
            if row_index % 2:
                set_cell_shading(cell, "F4F8F8")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            r.font.size = Pt(9.2)
            if widths:
                cell.width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.8)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.16
    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(34)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(14)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    for name, size, before, after in (("Heading 1", 22, 24, 10), ("Heading 2", 15, 18, 7), ("Heading 3", 12, 14, 5)):
        st = styles[name]
        st.font.name = "Arial"
        st.font.bold = True
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    for style_name in ("List Bullet", "List Number"):
        styles[style_name].font.name = "Arial"
        styles[style_name].font.size = Pt(10.5)


def build():
    arch = QA / "architecture.png"
    recovery = QA / "recovery.png"
    make_architecture_diagram(arch)
    make_recovery_diagram(recovery)

    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.78)
        section.right_margin = Inches(0.78)

    section = doc.sections[0]
    doc.settings.odd_and_even_pages_header_footer = True
    for footer in (section.footer, section.even_page_footer):
        footer_p = footer.paragraphs[0]
        footer_p.text = ""
        add_page_number(footer_p)

    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cover.paragraph_format.space_after = Pt(54)
    cover.add_run().add_picture(str(BRAND / "logo-horizontal.png"), width=Inches(4.8))
    p = doc.add_paragraph(style="Title")
    p.add_run("HearthMesh High Level Design")
    p_border = p._p.get_or_add_pPr().find(qn("w:pBdr"))
    if p_border is not None:
        p._p.get_or_add_pPr().remove(p_border)
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(34)
    r = sub.add_run("Global attributable emergency continuity mesh")
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor.from_string(MID)
    meta = [
        ("Document version", "0.1"),
        ("Project status", "Experimental pre-alpha"),
        ("Date", "24 September 2026"),
        ("Project owner", "Valerio Musella"),
    ]
    add_table(doc, ["Field", "Value"], meta, [1.8, 4.6])
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(42)
    r = p.add_run("Purpose")
    r.bold = True
    r.font.size = Pt(12)
    add_para(doc, "Define the target architecture and the deliberately narrower three-node MVP used to validate signed publication, replication, corruption detection and peer-assisted recovery.")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run("Status statement")
    r.bold = True
    r.font.size = Pt(12)
    add_para(doc, "This design is not a production security assessment. The MVP demonstrates selected system behavior and requires protocol hardening, independent review and multi-host testing before real continuity data is entrusted to it.")
    doc.add_page_break()

    add_heading(doc, "1 Executive summary", 1)
    add_para(doc, "HearthMesh is designed as a global, non-anonymous emergency continuity mesh. Independently operated nodes exchange readable public content that is signed, attributable, versioned and locally verified. A node can continue serving content already acquired when the original publisher, a peer or the wider Internet becomes unavailable.")
    add_para(doc, "The first implementation is intentionally smaller: three explicitly configured Go nodes on one Docker host. It proves the mechanics of signed HearthPacks, peer replication, local access, tamper detection and repair from a healthy replica. It does not yet provide global discovery, permissionless enrollment, private-cell confidentiality, key revocation or production-grade transport security.")
    add_para(doc, "The central architectural choice is accountability. The official design does not include onion routing, anonymous publication, payments, marketplaces or public encrypted blind storage. Runtime trust decisions use signatures, hashes and explicit operator policy; no AI model or inference service participates.")

    add_heading(doc, "2 Scope and design principles", 1)
    add_table(doc, ["Goal", "Architectural response"], [
        ("Continuity during disruption", "Serve locally complete packs without a central cloud or reachable publisher."),
        ("Attributable public publishing", "Bind each manifest to a persistent publisher key and visible operator record."),
        ("Integrity and provenance", "Sign immutable manifests and verify each referenced blob by SHA-256."),
        ("Failure recovery", "Pull a valid object from another approved peer after local corruption is detected."),
        ("Operator control", "Apply local publisher, category, size, retention and replication policies."),
        ("AI independence", "Use deterministic protocol rules with no model, agent or inference dependency."),
    ], [2.2, 4.2])
    add_heading(doc, "2.1 Explicit non-goals", 2)
    add_bullets(doc, [
        "Anonymous participation, hidden publishers, onion routing or traffic-obfuscation guarantees.",
        "A global consensus protocol or universally authoritative latest version.",
        "Guaranteed truth, legality, safety, completeness or availability of published information.",
        "Real-time emergency dispatch or replacement of official emergency channels.",
        "Arbitrary anonymous uploads, payments, marketplaces or public blind storage.",
        "Production private cells, automated moderation or secure key recovery in the MVP.",
    ])

    add_heading(doc, "3 Architecture overview", 1)
    doc.add_picture(str(arch), width=Inches(6.85))
    cap = doc.add_paragraph("Figure 1. Target architecture and the private-cell isolation boundary.")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(9)
    add_para(doc, "The target system uses the same core node software for the public mesh and future private cells, but the two spaces must not share discovery, indexes, membership or storage by default. For early private-cell work, separate processes and volumes are safer than unproven in-process multi-tenancy.")
    add_table(doc, ["Component", "Primary responsibility"], [
        ("Identity manager", "Loads node identity, publisher trust records and pinned peer identities."),
        ("Peer transport", "Connects directly to configured peers and applies timeouts and transfer limits."),
        ("Pack verifier", "Checks manifest structure, publisher signature, paths, media policy and blob hashes."),
        ("Content store", "Stores immutable manifests and content-addressed blobs with atomic promotion."),
        ("Replication scheduler", "Polls known peers and pulls locally eligible missing or damaged objects."),
        ("Integrity scanner", "Recomputes hashes, quarantines corrupt objects and requests repair."),
        ("Local portal", "Shows available packs, versions, provenance and node status without executing pack content."),
        ("Publisher tools", "Builds and signs HearthPacks through a local administrative path."),
    ], [1.75, 4.65])

    add_heading(doc, "4 Trust model and boundaries", 1)
    add_para(doc, "A valid signature proves control of a publisher key; it does not prove that the material is correct. An authenticated peer is identifiable; it is not automatically honest. HearthMesh therefore keeps publisher trust, peer admission and content verification as separate decisions.")
    add_table(doc, ["Boundary", "Risk", "Required control"], [
        ("Operator to local administration", "Policy deletion or misuse of signing material", "Local binding, restricted permissions and separated publisher keys."),
        ("Peer to transport", "False inventory, resource exhaustion or stale content", "Pinned identities, limits, pull replication and backoff."),
        ("Received bytes to storage", "Malformed, oversized or altered objects", "Staging, schema checks, signature verification and digest verification."),
        ("Publisher to reader", "Authentic but incorrect or harmful content", "Visible attribution, local trust tier and operator moderation."),
        ("Public mesh to private cell", "Metadata or content leakage", "Separate membership, namespace, storage, discovery and keys."),
        ("Host to node", "Key and data compromise", "Host hardening; acknowledge that host compromise defeats local guarantees."),
    ], [1.7, 2.1, 2.6])
    add_heading(doc, "4.1 Identity and trust tiers", 2)
    add_para(doc, "Node identity, publisher identity and accountable operator identity are distinct. An Ed25519 key alone creates a stable pseudonym, not verified real-world identity. Public enrollment must therefore attach an operator name or organization, contact route and explicit verification status to the cryptographic identity.")
    add_table(doc, ["Tier", "Meaning", "Default replication treatment"], [
        ("Blocked", "Locally prohibited publisher", "Reject and do not advertise."),
        ("Unrecognized", "Valid signature without operator approval", "Do not replicate automatically."),
        ("Enrolled", "Accountable publisher record approved locally", "Eligible under local content policy."),
        ("Endorsed", "Approved for a defined role or information domain", "Display endorsement and its exact scope."),
    ], [1.25, 2.55, 2.65])

    add_heading(doc, "5 HearthPack data model", 1)
    add_para(doc, "A HearthPack is an immutable signed manifest plus the immutable blobs referenced by that manifest. The pack identifier is derived from canonical manifest bytes; a blob identifier is derived from the exact blob bytes.")
    add_table(doc, ["Element", "Required content"], [
        ("Protocol metadata", "Protocol version, manifest schema version and algorithm identifiers."),
        ("Publisher", "Publisher identifier, public-key reference and signature envelope."),
        ("Publication identity", "Logical publication ID, revision and optional predecessor manifest digest."),
        ("Descriptive metadata", "Title, summary, language, topic tags, geography, license and review-by date."),
        ("Entries", "Normalized path, media type, byte size and SHA-256 digest for every file."),
        ("Classification", "Public in MVP; private classifications are rejected until private cells exist."),
    ], [1.75, 4.65])
    add_para(doc, "Protocol version 1 must freeze canonical encoding and signature inputs before interoperability is claimed. Ordinary JSON map serialization is not sufficient unless field ordering, whitespace, duplicate-field handling and numeric representations are specified and tested.")
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    r = p.add_run("blob_id = SHA-256(blob bytes)\npack_id = SHA-256(canonical manifest bytes)")
    r.font.name = "Courier New"
    r.font.size = Pt(10)

    add_heading(doc, "6 Networking and interfaces", 1)
    add_para(doc, "The MVP uses direct, explicitly configured peer addresses and periodic pull replication. The target design should use TLS 1.3 mutual authentication with peer identities pinned to expected node keys. Bootstrap peers provide initial addresses; they are not authorities.")
    add_table(doc, ["Interface", "Purpose", "Exposure"], [
        ("GET /v1/node", "Signed node descriptor and protocol capabilities", "Approved peers"),
        ("GET /v1/inventory", "Bounded pages of complete shareable pack IDs", "Approved peers"),
        ("GET /v1/manifests/{pack_id}", "Signed manifest envelope", "Approved peers"),
        ("GET /v1/blobs/{blob_id}", "Verified immutable blob bytes", "Approved peers"),
        ("Local portal", "Browse, verify and download installed public material", "Local network by policy"),
        ("Administrative API", "Publish, trust, retention and maintenance actions", "Local host only"),
    ], [1.8, 3.0, 1.6])
    add_para(doc, "The target protocol requires bounded response sizes, timeouts, concurrency limits, retry backoff and jitter. Nodes must never dial arbitrary endpoints embedded in a manifest or automatically admit transitively advertised peers.")

    add_heading(doc, "7 Core system flows", 1)
    add_heading(doc, "7.1 Publish", 2)
    add_numbered(doc, [
        "The local operator selects files and publication metadata.",
        "Publisher tooling validates paths, sizes and allowed media types.",
        "Files are hashed and a canonical manifest is signed.",
        "The local node independently verifies the pack and publisher policy.",
        "Objects move from staging into immutable storage only after verification.",
        "A complete pack becomes visible locally and eligible for peer inventory.",
    ])
    add_heading(doc, "7.2 Replicate", 2)
    add_numbered(doc, [
        "A node polls an approved peer inventory.",
        "It fetches an unfamiliar signed manifest and evaluates local admission policy.",
        "Missing blobs are pulled into bounded staging storage.",
        "Every size and digest is verified before the pack is committed.",
        "Only complete packs appear in the local portal or outgoing inventory.",
    ])
    add_heading(doc, "7.3 Detect and repair corruption", 2)
    doc.add_picture(str(recovery), width=Inches(6.85))
    cap = doc.add_paragraph("Figure 2. The MVP recovery case: publisher loss, damaged replica and verified repair.")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(9)
    add_para(doc, "On a digest mismatch, the object is quarantined and dependent packs become unavailable. The scheduler asks other configured peers for the expected digest, verifies replacement bytes and installs them atomically. Hashes detect damage but cannot reconstruct data; recovery fails if every correct replica has been lost.")

    add_heading(doc, "8 Public mesh and private continuity cells", 1)
    add_table(doc, ["Property", "Global public mesh", "Private continuity cell"], [
        ("Purpose", "Broadly distributable emergency and continuity knowledge", "Restricted operational or personal material"),
        ("Content visibility", "Readable public content", "Authorized members only"),
        ("Publisher model", "Attributable publisher with local trust status", "Known membership and cell-scoped authorization"),
        ("Discovery", "Future global federation; allowlisted in MVP", "Isolated and not publicly advertised"),
        ("Storage", "Public immutable objects", "Separate encrypted store and index"),
        ("MVP status", "Partially demonstrated", "Architecture placeholder only"),
    ], [1.45, 2.45, 2.45])
    add_para(doc, "A private cell is not a hidden anonymous channel. It requires explicit membership, isolated peer lists, authenticated encrypted transport, data-at-rest policy and tested key management. Revoking membership cannot erase copies already obtained.")

    add_heading(doc, "9 Security and misuse resistance", 1)
    add_para(doc, "The MVP should demonstrate cryptographic integrity, not broad security. It must reject malformed manifests, unsafe paths, altered blobs and unsupported algorithms. Publisher-supplied code must never execute in the portal, and the node must not automatically extract archives or fetch embedded external resources.")
    add_table(doc, ["Threat", "MVP response", "Remaining exposure"], [
        ("Dishonest publisher", "Visible signature and local trust policy", "Signature cannot establish truth."),
        ("Compromised publisher key", "Manual block configuration", "Disconnected nodes may not learn the revocation."),
        ("Malicious peer", "Allowlist, pull transfer and limits", "Approved peer can suppress or waste resources."),
        ("Corrupt local object", "Digest scan, quarantine and peer repair", "No repair if all valid copies are lost."),
        ("Path traversal", "Normalize and reject unsafe entry paths", "Parser defects still require testing and fuzzing."),
        ("Active file content", "MVP allows a small passive format set", "A valid file may still exploit an external viewer."),
        ("Sybil enrollment", "Manual enrollment in MVP", "Global admission remains unsolved."),
        ("Traffic analysis", "No concealment is attempted", "Peer relationships and timing remain visible."),
    ], [1.55, 2.35, 2.45])
    add_heading(doc, "9.1 Misuse constraints", 2)
    add_bullets(doc, [
        "No onion routing, hidden services, anonymous publisher mode or traffic-obfuscation claim.",
        "No public encrypted blind-storage service in the official network.",
        "No payments, marketplace, anonymous messaging or arbitrary unauthenticated upload endpoint.",
        "Public packs remain readable and attributable; node operators choose what they replicate.",
        "Moderation and deletion are local. Global recall of replicated material cannot be guaranteed.",
    ])
    add_para(doc, "These choices reduce darknet-like utility but cannot prevent modification or misuse of open-source code. The project should describe enforceable protocol behavior precisely and avoid claiming that file-extension checks can prove the absence of hidden or encrypted data.")

    add_heading(doc, "10 Failure behavior", 1)
    add_table(doc, ["Failure", "Required behavior"], [
        ("One peer unavailable", "Continue local service and retry another approved peer."),
        ("All peers unavailable", "Serve complete local packs and show disconnected status."),
        ("Interrupted transfer", "Keep bounded staging state; never expose an incomplete pack."),
        ("Digest mismatch", "Quarantine the object, record the event and request another source."),
        ("Invalid signature", "Reject the manifest with a visible reason."),
        ("Disk full", "Stop acquisition safely and preserve installed content where possible."),
        ("Conflicting revisions", "Retain and show both branches; do not resolve silently by timestamp."),
        ("Lost node key", "Require explicit reenrollment under a new identity."),
        ("All correct copies lost", "Declare content unavailable; do not imply recoverability."),
        ("MVP Docker host lost", "All three logical nodes may fail because they share one physical host."),
    ], [2.05, 4.35])

    add_heading(doc, "11 Deployment and operations", 1)
    add_para(doc, "Docker Compose starts three logical Go nodes with distinct identities, ports, configurations and persistent volumes. This is reproducible and sufficient for behavior testing, but all nodes share one host, disk, power source and network path. A meaningful resilience pilot requires independent machines and failure domains.")
    add_table(doc, ["Operational signal", "Interpretation"], [
        ("Process healthy", "The node process is running and local checks pass."),
        ("Content ready", "At least the required locally pinned packs are complete and verified."),
        ("Mesh degraded", "One or more peers are unavailable; local continuity may still be intact."),
        ("Pack quarantined", "A signature, schema or digest verification failed."),
        ("Storage pressure", "Acquisition may stop before installed pinned content is affected."),
    ], [2.0, 4.4])
    add_para(doc, "Logs and metrics should cover peer reachability, pack states, replication retries, bytes transferred, verification failures, integrity-scan coverage, disk quotas and protocol versions. Private content and unnecessary reader identifiers must not be logged.")

    add_heading(doc, "12 MVP implementation", 1)
    add_table(doc, ["Capability", "MVP position"], [
        ("Runtime", "Single Go binary per node"),
        ("Topology", "Three allowlisted nodes with explicit bootstrap peers"),
        ("Identity", "Persistent Ed25519 node and publisher material"),
        ("Integrity", "Signed manifest plus SHA-256 file verification"),
        ("Replication", "Direct peer polling and pull transfer"),
        ("Recovery", "Detect modified bytes and obtain a valid peer copy"),
        ("Access", "Local read-only web view and status endpoints"),
        ("Packaging", "Docker Compose and scripted demonstration"),
        ("Status", "Experimental pre-alpha; not production-secure"),
    ], [2.0, 4.4])
    add_heading(doc, "12.1 Demonstration acceptance", 2)
    add_numbered(doc, [
        "Node A publishes a signed sample HearthPack.",
        "Nodes B and C replicate and verify the same pack.",
        "Node A is stopped and no longer serves as publisher.",
        "A stored file on node B is deliberately modified.",
        "Node B detects that the file digest does not match the signed manifest.",
        "Node B retrieves the expected bytes from node C and verifies them before restoration.",
        "The recovered pack remains available through the local portal.",
    ])
    add_para(doc, "Acceptance was completed on 24 September 2026 on a Windows x86-64 host. All four Go tests passed, the Windows binary built successfully, and the three-node Docker drill demonstrated publication, replication, continued availability after node A stopped, corruption detection on node B and repair from node C. This is single-host functional evidence, not an independent security audit or a multi-host resilience test. Continuous integration must repeat tests and builds for ongoing changes.")

    add_heading(doc, "13 Architecture decisions", 1)
    add_table(doc, ["Decision", "Rationale"], [
        ("Immutable content addressing", "Makes stored objects independently verifiable and avoids silent mutation."),
        ("Separate publisher and node identities", "Publishing authority should not be implied by transport participation."),
        ("Direct identifiable peers", "Supports accountability and intentionally rejects anonymity infrastructure."),
        ("Pull replication", "Prevents unsolicited peers from pushing arbitrary content into storage."),
        ("Local policy", "Signing does not create a storage obligation for every node."),
        ("No global consensus", "The system can preserve conflicting branches without pretending to select truth."),
        ("No AI at runtime", "Trust decisions remain deterministic, inspectable and service-independent."),
        ("Separate private-cell process", "Avoids claiming safe multi-tenancy before isolation is designed and tested."),
    ], [2.1, 4.3])

    add_heading(doc, "14 Roadmap", 1)
    add_table(doc, ["Phase", "Outcome", "Exit evidence"], [
        ("1 Protocol specification", "Canonical encoding, IDs, limits and test vectors", "Independent implementations produce identical IDs and signatures."),
        ("2 Three-node MVP", "Publication, replication, portal, tamper detection and repair", "Automated tests plus a recorded repeatable Docker drill."),
        ("3 Multi-host pilot", "Independent operators and physical failure domains", "Measured recovery during network and publisher loss."),
        ("4 Security hardening", "Rotation, revocation, fuzzing, abuse response and dependency review", "Independent assessment and tracked remediation."),
        ("5 Private cells", "Membership, isolation, encryption and key lifecycle", "Confidentiality and revocation-limit tests."),
        ("6 Broader federation", "Scalable discovery and inventory exchange", "Pilot evidence without a mandatory central coordinator."),
    ], [1.4, 2.5, 2.5])

    add_heading(doc, "15 Principal limitations", 1)
    add_bullets(doc, [
        "The MVP proves a local three-node behavior, not an operational global mesh.",
        "Three containers on one machine do not provide hardware, power or geographic resilience.",
        "A valid signature proves key control and byte integrity, not factual correctness.",
        "Manual enrollment reduces the initial attack surface but does not solve global identity or Sybil resistance.",
        "Key rotation, revocation, recovery and publisher withdrawal require protocol and governance work.",
        "Public replication makes universal deletion impossible once another party retains a copy.",
        "Private continuity cells are an architectural direction, not an implemented confidentiality feature.",
        "Security claims require independent review, adversarial testing and operational evidence.",
    ])
    add_para(doc, "The correct pre-alpha claim is narrow: HearthMesh demonstrates how attributable signed content can remain locally available, detect alteration and recover from another known peer without relying on the original publisher.")

    output = OUT / "HearthMesh_High_Level_Design_v0.1.docx"
    doc.save(output)
    print(output)


if __name__ == "__main__":
    build()
